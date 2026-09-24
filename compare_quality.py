"""
品質比較スクリプト：同じお題を「モデル × 品質」の組み合わせで生成し、比較用の PowerPoint にまとめる。

使い方（リポジトリ直下で）:
  .venv/bin/python compare_quality.py --dry-run     # 生成せずに計画（枚数）だけ表示
  .venv/bin/python compare_quality.py               # 既定: 3お題 × 6条件 = 18枚（実行前に確認あり。-y で省略）
  .venv/bin/python compare_quality.py --prompts text,diagram
  .venv/bin/python compare_quality.py --variants 2:high,flare:high,sunburst:high
  .venv/bin/python compare_quality.py --prompt "自由なお題"    # 任意のお題1つで比較
  .venv/bin/python compare_quality.py --resume output/compare_YYYYMMDD_HHMMSS
      # 途中で止まった回の続きを生成（成功済みはスキップ）。全部済みならスライドだけ作り直す

出力: output/compare_YYYYMMDD_HHMMSS/ に画像・results.json・quality_compare.pptx
生成時間を公平に比べるため、API 呼び出しは1枚ずつ順番に行う。
"""

import argparse
import base64
import json
import math
import sys
import time
from datetime import datetime
from pathlib import Path

from openai import OpenAI
from PIL import Image

import generate as g

MODEL_ALIASES = {
    "2": "gpt-image-2",
    "flare": "gpt-image-2.5-flare",
    "sunburst": "gpt-image-2.5-sunburst",
}

# 旧既定（gpt-image-2 / high）を基準に、flare の全品質を並べる
DEFAULT_VARIANTS = ["2:high", "flare:low", "flare:medium", "flare:high", "flare:xhigh", "flare:max"]

# スライド用途で差が出やすい3種（日本語の文字・図解の細部・写真の質感）
PROMPTS = {
    "text": (
        "日本語タイトル",
        "ビジネスプレゼンテーションの表紙スライド用キービジュアル。画面中央に大きく太いゴシック体で"
        "「2026年度 事業戦略」、その下に一回り小さく「成長・効率化・人材育成」と正確に描画する。"
        "背景は深い紺色から青へのグラデーションに、抽象的な光のラインと粒子。"
        "文字は指定したもの以外は入れない。",
    ),
    "diagram": (
        "図解（4ステップ）",
        "スライド用のフラットデザインの図解。白い背景に、左から右へ矢印でつながる4つの角丸カード。"
        "各カードにシンプルなアイコンと日本語ラベル「企画」「設計」「開発」「運用」を順に配置し、"
        "各ラベルの下に小さく「STEP 1」〜「STEP 4」。配色は紺とオレンジのアクセント。"
        "文字は指定したもの以外は入れない。",
    ),
    "photo": (
        "写真リアル",
        "明るくモダンなオフィスで、ホワイトボードの前でディスカッションする4人のビジネスパーソン。"
        "自然光、浅い被写界深度、写真のようにリアル。人物は画面の左側に寄せ、右側3分の1は柔らかく"
        "ぼけた壁で文字を載せられる余白にする。手や顔の描写は自然に。",
    ),
}

RESULTS_FILE = "results.json"
DECK_FILE = "quality_compare.pptx"


def _parse_variant(spec: str) -> dict:
    model, sep, quality = spec.strip().rpartition(":")
    if not sep or not model or not quality:
        raise ValueError(f"条件 '{spec}' は 'モデル:品質' の形で指定してください（例: flare:high）")
    return {"model": MODEL_ALIASES.get(model, model), "quality": quality}


def _variant_label(v: dict) -> str:
    return f"{v['model']} / {v['quality']}"


def _result_key(prompt_key: str, v: dict) -> str:
    return f"{prompt_key}__{v['model']}__{v['quality']}"


def _text_width_cm(text: str, pt: float) -> float:
    """1行テキストのおおよその幅（全角=1em、半角=0.55em）。"""
    em = pt * 2.54 / 72
    return sum(em if ord(ch) > 0x2E7F else em * 0.55 for ch in text)


def _fmt_bytes(n: float) -> str:
    return f"{n / 1e6:.1f} MB" if n >= 1e5 else f"{n / 1e3:.0f} KB"


def _save_state(run_dir: Path, state: dict) -> None:
    tmp = run_dir / (RESULTS_FILE + ".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(run_dir / RESULTS_FILE)


def _generate_one(client: OpenAI, prompt: str, v: dict, size: str) -> tuple[bytes, dict]:
    """1枚生成して (PNGバイト列, 計測値) を返す。"""
    started = time.monotonic()
    response = client.images.generate(
        model=v["model"], prompt=prompt, size=size, quality=v["quality"], n=1,
    )
    elapsed = time.monotonic() - started
    if not response.data or response.data[0].b64_json is None:
        raise RuntimeError("APIレスポンスに画像データがありません")
    usage = getattr(response, "usage", None)
    return base64.b64decode(response.data[0].b64_json), {
        "elapsed_s": round(elapsed, 1),
        "api_quality": getattr(response, "quality", None),
        "usage": {
            "input_tokens": usage.input_tokens,
            "output_tokens": usage.output_tokens,
        } if usage else None,
    }


def run(state: dict, run_dir: Path, timeout: float) -> None:
    done = {r["key"] for r in state["results"] if r["ok"]}
    todo = [(p, v) for p in state["prompts"] for v in state["variants"]
            if _result_key(p["key"], v) not in done]
    print(f"生成 {len(todo)} 枚（済み {len(done)} 枚）→ {run_dir}")
    if not todo:
        return
    client = OpenAI(api_key=g._require_api_key(), timeout=timeout, max_retries=g.MAX_RETRIES)

    for i, (p, v) in enumerate(todo, 1):
        key = _result_key(p["key"], v)
        print(f"[{i}/{len(todo)}] {p['label']} × {_variant_label(v)} ...", flush=True)
        entry = {"key": key, "prompt_key": p["key"], **v, "ok": False}
        try:
            png, metrics = _generate_one(client, p["prompt"], v, state["resolved_size"])
            path = run_dir / f"{key}.png"
            path.write_bytes(png)
            with Image.open(path) as im:
                px = f"{im.size[0]}x{im.size[1]}"
            entry.update(ok=True, file=path.name, px=px, bytes=len(png), **metrics)
            out_tok = (metrics["usage"] or {}).get("output_tokens")
            print(f"    {metrics['elapsed_s']}秒 / {px}" + (f" / 出力 {out_tok:,} tok" if out_tok else ""))
        except Exception as e:  # noqa: BLE001 1条件の失敗で比較全体を止めない
            entry["error"] = f"{type(e).__name__}: {e}"
            print(f"    失敗: {entry['error']}")
        state["results"] = [r for r in state["results"] if r["key"] != key] + [entry]
        _save_state(run_dir, state)


# ---------------------------------------------------------------- スライド作成

SLIDE_W_CM, SLIDE_H_CM = 30.0, 18.2  # 普段のスライド（30cm×18.2cm）に合わせる
DARK = "16181D"
PANEL = "2A2E36"
WHITE = "FFFFFF"
MUTED = "A3A9B5"
INK = "1F2328"
ACCENT = "F2A93B"  # 基準（旧既定）の目印
FONT_LATIN = "Arial"
FONT_EA = "Yu Gothic"


def build_deck(state: dict, run_dir: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Cm, Pt
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(SLIDE_W_CM), Cm(SLIDE_H_CM)
    blank = prs.slide_layouts[6]
    results = {r["key"]: r for r in state["results"]}
    baseline = state["variants"][0]
    w_px, h_px = (int(x) for x in state["resolved_size"].split("x"))
    aspect = w_px / h_px

    def new_slide(bg: str):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        return s

    def style_run(run, size, color, bold=False):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.name = FONT_LATIN
        rpr = run._r.get_or_add_rPr()
        ea = etree.SubElement(rpr, qn("a:ea"))
        ea.set("typeface", FONT_EA)
        rpr.find(qn("a:latin")).addnext(ea)

    def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        """lines: [(text, size, color, bold[, space_before_pt]), ...] を1段落ずつ置く。"""
        tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (text, size, color, bold, *space) in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            if space:
                para.space_before = Pt(space[0])
            run = para.add_run()
            run.text = text
            style_run(run, size, color, bold)
        return tb

    def add_box(slide, x, y, w, h, color, alpha=None, rounded=False):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Cm(x), Cm(y), Cm(w), Cm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        if alpha is not None:
            clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def add_image_fit(slide, path, x, y, w, h):
        """枠 (x, y, w, h) に縦横比を保って中央配置する。"""
        with Image.open(path) as im:
            ratio = im.size[0] / im.size[1]
        iw, ih = (w, w / ratio) if w / h < ratio else (h * ratio, h)
        slide.shapes.add_picture(str(path), Cm(x + (w - iw) / 2), Cm(y + (h - ih) / 2), Cm(iw), Cm(ih))

    def metrics_line(r: dict) -> str:
        parts = [f"{r['elapsed_s']}秒"]
        out_tok = (r.get("usage") or {}).get("output_tokens")
        if out_tok:
            parts.append(f"出力 {out_tok:,} tok")
        parts.append(r["px"].replace("x", "×"))
        return " · ".join(parts)

    def set_cell(cell, text, size, color, bold=False, fill=None, align=PP_ALIGN.LEFT):
        cell.text = ""
        para = cell.text_frame.paragraphs[0]
        para.alignment = align
        run = para.add_run()
        run.text = text
        style_run(run, size, color, bold)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Cm(0.25)
        if fill:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)

    def add_bar_chart(slide, x, y, w, h, title, labels, values, number_format):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        data = CategoryChartData()
        data.categories = labels
        data.add_series(title, values)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, Cm(x), Cm(y), Cm(w), Cm(h), data).chart
        chart.has_legend = False
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        style_run(chart.chart_title.text_frame.paragraphs[0].runs[0], 12, INK, True)
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        plot.data_labels.number_format = number_format
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size = Pt(10)
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string("3B4A6B")
        series.points[0].format.fill.solid()  # 先頭＝基準をアクセント色に
        series.points[0].format.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        chart.category_axis.reverse_order = True  # 表と同じ上から順に並べる
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.format.line.color.rgb = RGBColor.from_string("D1D5DB")
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False

    # 表紙
    s = new_slide(DARK)
    add_text(s, 1.5, 1.4, 27, 2.0, [("GPT Image 品質比較", 36, WHITE, True)])
    add_text(s, 1.5, 3.4, 27, 1.0, [(
        f"{state['created'][:10]} ・ サイズ {state['size']}（{w_px}×{h_px}）・ "
        f"{len(state['prompts'])} お題 × {len(state['variants'])} 条件", 14, MUTED, False)])
    rows = len(state["variants"]) + 1
    tbl = s.shapes.add_table(rows, 2, Cm(1.5), Cm(5.2), Cm(11.5), Cm(0.9 * rows)).table
    tbl.columns[0].width, tbl.columns[1].width = Cm(7.5), Cm(4.0)
    set_cell(tbl.cell(0, 0), "モデル", 12, WHITE, True, PANEL)
    set_cell(tbl.cell(0, 1), "品質", 12, WHITE, True, PANEL)
    for i, v in enumerate(state["variants"], 1):
        color = ACCENT if v == baseline else WHITE
        set_cell(tbl.cell(i, 0), v["model"] + ("（基準）" if v == baseline else ""), 12, color, False, DARK)
        set_cell(tbl.cell(i, 1), v["quality"], 12, color, False, DARK)
    prompt_lines = [("お題", 14, WHITE, True)]
    for p in state["prompts"]:
        prompt_lines += [(p["label"], 12, WHITE, True, 10),
                         (p["prompt"][:80] + ("…" if len(p["prompt"]) > 80 else ""), 10, MUTED, False, 2)]
    add_text(s, 14.5, 5.2, 14, 10.5, prompt_lines)
    add_text(s, 1.5, 16.2, 27, 0.8, [(
        "見方：お題ごとに「一覧」→「条件ごとの全面表示（実際の貼り付けサイズ）」。最後に数値まとめ。",
        11, MUTED, False)])

    # お題ごと：一覧 → 全面表示
    n = len(state["variants"])
    cols = n if n <= 3 else 3 if n <= 6 else 4
    grid_rows = math.ceil(n / cols)
    gap, caption_h, top, margin_x, bottom = 0.5, 1.2, 3.0, 1.0, 0.6
    cell_w = (SLIDE_W_CM - 2 * margin_x - gap * (cols - 1)) / cols
    img_h = cell_w / aspect
    avail_h = SLIDE_H_CM - top - bottom - gap * (grid_rows - 1)
    if grid_rows * (img_h + caption_h) > avail_h:
        img_h = avail_h / grid_rows - caption_h
        cell_w = img_h * aspect
    x0 = (SLIDE_W_CM - (cell_w * cols + gap * (cols - 1))) / 2

    for p in state["prompts"]:
        s = new_slide(DARK)
        add_text(s, margin_x, 0.5, SLIDE_W_CM - 2 * margin_x, 1.1, [(f"お題：{p['label']}", 20, WHITE, True)])
        add_text(s, margin_x, 1.7, SLIDE_W_CM - 2 * margin_x, 1.0, [
            (p["prompt"][:160] + ("…" if len(p["prompt"]) > 160 else ""), 9, MUTED, False)])
        for i, v in enumerate(state["variants"]):
            x = x0 + (i % cols) * (cell_w + gap)
            y = top + (i // cols) * (img_h + caption_h + gap)
            r = results.get(_result_key(p["key"], v))
            is_base = v == baseline
            label = _variant_label(v) + ("（基準）" if is_base else "")
            if r and r["ok"]:
                add_image_fit(s, run_dir / r["file"], x, y, cell_w, img_h)
                sub = metrics_line(r)
            else:
                add_box(s, x, y, cell_w, img_h, PANEL)
                err = (r or {}).get("error", "未生成")
                add_text(s, x + 0.4, y + 0.4, cell_w - 0.8, img_h - 0.8,
                         [("生成失敗" if r else "未生成", 14, WHITE, True), (err[:120], 9, MUTED, False)],
                         anchor=MSO_ANCHOR.MIDDLE)
                sub = "—"
            add_text(s, x, y + img_h + 0.15, cell_w, caption_h - 0.15, [
                (label, 11, ACCENT if is_base else WHITE, True), (sub, 9, MUTED, False)])

        for v in state["variants"]:
            r = results.get(_result_key(p["key"], v))
            if not (r and r["ok"]):
                continue
            s = new_slide("000000")
            add_image_fit(s, run_dir / r["file"], 0, 0, SLIDE_W_CM, SLIDE_H_CM)
            badge = f"{p['label']}｜{_variant_label(v)}｜{metrics_line(r)}"
            badge_w = _text_width_cm(badge, 11) * 1.1 + 0.4
            add_box(s, 0.4, 0.4, badge_w + 0.8, 0.9, "000000", alpha=0.6, rounded=True)
            add_text(s, 0.8, 0.4, badge_w, 0.9, [(badge, 11, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
            usage = r.get("usage") or {}
            s.notes_slide.notes_text_frame.text = "\n".join([
                f"モデル: {v['model']} / 品質: {v['quality']}（API応答の品質: {r.get('api_quality') or '-'}）",
                f"生成時間: {r['elapsed_s']}秒 / 実寸: {r['px']} / ファイル: {r['file']}",
                f"トークン: 入力 {usage.get('input_tokens', '-')} / 出力 {usage.get('output_tokens', '-')}",
                "", "prompt:", p["prompt"],
            ])

    # 数値まとめ（表＋グラフ）
    stats = []
    for v in state["variants"]:
        rs = [results.get(_result_key(p["key"], v)) for p in state["prompts"]]
        ok = [r for r in rs if r and r["ok"]]
        toks = [r["usage"]["output_tokens"] for r in ok if r.get("usage")]
        stats.append({
            "v": v, "tried": len(rs), "ok": len(ok),
            "time": sum(r["elapsed_s"] for r in ok) / len(ok) if ok else None,
            "tok": sum(toks) / len(toks) if toks else None,
            "bytes": sum(r["bytes"] for r in ok) / len(ok) if ok else None,
        })

    s = new_slide(WHITE)
    add_text(s, 1.5, 0.9, 27, 1.5, [("数値まとめ", 28, INK, True)])
    headers = ["条件", "生成時間", "出力トークン", "ファイルサイズ", "成功", "所感（記入欄）"]
    widths = [8.0, 3.2, 3.6, 3.6, 1.8, 6.8]
    row_h = 0.85
    tbl = s.shapes.add_table(n + 1, len(headers), Cm(1.5), Cm(2.8), Cm(sum(widths)),
                             Cm(row_h * (n + 1))).table
    for c, (h, w) in enumerate(zip(headers, widths)):
        tbl.columns[c].width = Cm(w)
        set_cell(tbl.cell(0, c), h, 12, WHITE, True, INK)
    for i, st in enumerate(stats, 1):
        is_base = st["v"] == baseline
        cells = [
            _variant_label(st["v"]) + ("（基準）" if is_base else ""),
            f"{st['time']:.1f} 秒" if st["time"] is not None else "—",
            f"{st['tok']:,.0f}" if st["tok"] is not None else "—",
            _fmt_bytes(st["bytes"]) if st["bytes"] is not None else "—",
            f"{st['ok']}/{st['tried']}",
            "",
        ]
        fill = "FDF3E1" if is_base else ("F4F5F7" if i % 2 == 0 else WHITE)
        for c, text in enumerate(cells):
            set_cell(tbl.cell(i, c), text, 12, INK, c == 0, fill,
                     PP_ALIGN.LEFT if c in (0, 5) else PP_ALIGN.RIGHT)
        tbl.rows[i].height = Cm(row_h)
    tbl.rows[0].height = Cm(row_h)

    # グラフは表の下に収まれば同じスライド、収まらなければ次のスライドへ
    chart_y = 2.8 + row_h * (n + 1) + 0.8
    if chart_y + 5.0 > 16.0:
        s = new_slide(WHITE)
        add_text(s, 1.5, 0.9, 27, 1.5, [("数値まとめ（グラフ）", 28, INK, True)])
        chart_y = 2.8
    chart_h = 16.0 - chart_y
    labels = [_variant_label(st["v"]) for st in stats]
    add_bar_chart(s, 1.5, chart_y, 13.2, chart_h, "平均生成時間（秒）", labels,
                  [st["time"] for st in stats], "0.0")
    if any(st["tok"] is not None for st in stats):
        add_bar_chart(s, 15.3, chart_y, 13.2, chart_h, "平均出力トークン（料金の目安）", labels,
                      [st["tok"] for st in stats], "#,##0")
    add_text(s, 1.5, 16.5, 27, 0.8, [(
        "数値はお題ごとの平均。出力トークンは API が usage を返した場合のみ。生成時間はリトライを含む実測値。",
        10, "6B7280", False)])

    path = run_dir / DECK_FILE
    prs.save(path)
    return path


def main() -> None:
    ap = argparse.ArgumentParser(description="モデル × 品質の比較スライドを作る")
    ap.add_argument("--prompts", default=",".join(PROMPTS),
                    help=f"使うお題（カンマ区切り）: {', '.join(PROMPTS)}")
    ap.add_argument("--prompt", help="任意のお題（指定時は --prompts の代わりにこれ1つで比較）")
    ap.add_argument("--variants", default=",".join(DEFAULT_VARIANTS),
                    help="比較条件 'モデル:品質' のカンマ区切り。先頭が基準。"
                         f"モデル略称: {', '.join(f'{k}={v}' for k, v in MODEL_ALIASES.items())}")
    ap.add_argument("--size", default="slide", help="サイズ（generate.py のプリセット名か WxH）。既定 slide")
    ap.add_argument("--output-dir", default=str(g.OUTPUT_DIR), help="出力先の親ディレクトリ")
    ap.add_argument("--timeout", type=float, default=600, help="1枚あたりのタイムアウト秒（既定600）")
    ap.add_argument("--resume", help="既存の比較ディレクトリを指定して続きから実行（お題・条件は前回のものを使う）")
    ap.add_argument("--dry-run", action="store_true", help="生成せずに計画だけ表示")
    ap.add_argument("--yes", "-y", action="store_true", help="実行前の確認を省略")
    args = ap.parse_args()

    if args.resume:
        run_dir = Path(args.resume).expanduser()
        state = json.loads((run_dir / RESULTS_FILE).read_text(encoding="utf-8"))
    else:
        if args.prompt:
            prompts = [{"key": "custom", "label": "任意のお題", "prompt": args.prompt}]
        else:
            keys = [k.strip() for k in args.prompts.split(",") if k.strip()]
            unknown = [k for k in keys if k not in PROMPTS]
            if unknown:
                sys.exit(f"不明なお題: {unknown}（選べるのは {list(PROMPTS)}）")
            prompts = [{"key": k, "label": PROMPTS[k][0], "prompt": PROMPTS[k][1]} for k in keys]
        variants = []
        try:
            for spec in args.variants.split(","):
                if spec.strip() and (v := _parse_variant(spec)) not in variants:
                    variants.append(v)
            resolved = {g._resolve_size(args.size, v["model"]) for v in variants}
            for v in variants:  # 課金前に不正な組み合わせ（例: 2:max）を弾く
                g._validate_quality(v["quality"], v["model"])
        except ValueError as e:
            sys.exit(f"エラー: {e}")
        if len(resolved) != 1 or "auto" in resolved:
            sys.exit(f"全条件で同じ実サイズ（WxH）になるサイズを指定してください: {sorted(resolved)}")
        state = {
            "created": datetime.now().isoformat(timespec="seconds"),
            "size": args.size,
            "resolved_size": resolved.pop(),
            "prompts": prompts,
            "variants": variants,
            "results": [],
        }
        run_dir = Path(args.output_dir).expanduser() / f"compare_{datetime.now():%Y%m%d_%H%M%S}"

    total = len(state["prompts"]) * len(state["variants"])
    print(f"お題 {len(state['prompts'])} × 条件 {len(state['variants'])} = {total} 枚"
          f"（{state['resolved_size']}）")
    for v in state["variants"]:
        print(f"  - {v['model']} / {v['quality']}")
    if args.dry_run:
        return
    if not args.yes and input("API で生成します（料金が発生します）。続行しますか？ [y/N] ").strip().lower() != "y":
        print("中止しました")
        return

    run_dir.mkdir(parents=True, exist_ok=True)
    _save_state(run_dir, state)
    run(state, run_dir, args.timeout)
    deck = build_deck(state, run_dir)
    print(f"\nスライド: {deck}")


if __name__ == "__main__":
    main()
